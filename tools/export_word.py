#!/usr/bin/env python3
"""Export REPORT_WORD.md to editable DOCX. Requires python-docx and Pillow.

Usage: python3 tools/export_word.py
Legacy DOC can be produced with LibreOffice's MS Word 97 filter.
"""
import re
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.opc.constants import RELATIONSHIP_TYPE

ROOT = Path(__file__).resolve().parents[1]


def inline(paragraph, text):
    """Preserve bold, inline code and clickable external links."""
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`|\[[^]]+\]\(https?://[^)]+\))', text):
        if not part:
            continue
        match = re.fullmatch(r'\[([^]]+)\]\((https?://[^)]+)\)', part)
        if match:
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(match[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '245B83'); props.append(color)
            underline = OxmlElement('w:u'); underline.set(qn('w:val'), 'single'); props.append(underline)
            run.append(props)
            node = OxmlElement('w:t'); node.text = match[1]; run.append(node)
            link.append(run); paragraph._p.append(link)
        elif part.startswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('`'):
            run = paragraph.add_run(part[1:-1]); run.font.color.rgb = RGBColor.from_string('245B83')
        else:
            paragraph.add_run(part)


def table(document, rows):
    tbl = document.add_table(rows=0, cols=len(rows[0]))
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = {2: [6.3, 11.1], 3: [5.4, 6.4, 5.6], 5: [4.5, 2.5, 3.5, 3.4, 3.5]}[len(rows[0])]
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Cm(width)
    for index, row in enumerate(rows):
        cells = tbl.add_row().cells
        for cell, value, width in zip(cells, row, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.line_spacing = 1
            inline(p, value)
            for run in p.runs:
                run.font.size = Pt(9.5 if len(rows[0]) < 5 else 9)
                if index == 0: run.bold = True
            if index == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E6EEF5'); cell._tc.get_or_add_tcPr().append(shade)
        no_split = OxmlElement('w:cantSplit'); tbl.rows[index]._tr.get_or_add_trPr().append(no_split)
        if index == 0:
            repeat = OxmlElement('w:tblHeader'); tbl.rows[index]._tr.get_or_add_trPr().append(repeat)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def main():
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = Cm(1.5); section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.8); section.right_margin = Cm(1.8)
    section.header_distance = Cm(.7); section.footer_distance = Cm(.7)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'List Bullet', 'Caption']:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Times New Roman')
        style.font.size = Pt(11)
    normal = document.styles['Normal'].paragraph_format
    normal.space_after = Pt(6); normal.line_spacing = 1.06
    document.styles['Title'].font.size = Pt(20)
    document.styles['Title'].font.bold = True
    document.styles['Title'].font.color.rgb = RGBColor.from_string('173D5D')
    for name in ['Heading 1', 'Heading 2']:
        style = document.styles[name]; style.font.size = Pt(13); style.font.bold = True
        style.font.color.rgb = RGBColor.from_string('173D5D')
        style.paragraph_format.space_before = Pt(8); style.paragraph_format.space_after = Pt(5)
        style.paragraph_format.keep_with_next = True
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'vi-VN')
    document.styles['Normal']._element.get_or_add_rPr().append(lang)
    header = section.header.paragraphs[0]
    header.add_run('ROS 2 HUMBLE  •  TURTLEBOT 4  •  SLAM & NAV2').font.size = Pt(8)
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Báo cáo thực hành TurtleBot 4  |  Trang ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)

    source = ROOT / 'docs/REPORT_WORD.md'
    lines = source.read_text().splitlines(); i = 0; pending = []
    def flush():
        if pending:
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            inline(p, ' '.join(pending)); pending.clear()
    while i < len(lines):
        line = lines[i]
        if line == '<!-- PAGEBREAK -->':
            flush(); document.add_page_break()
        elif line.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [x.strip() for x in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r'[:\- ]+', x) for x in cells): rows.append(cells)
                i += 1
            table(document, rows); continue
        elif line.startswith('# '):
            flush(); p = document.add_paragraph(line[2:], 'Title'); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith('## '):
            flush(); document.add_heading(line[3:], level=1)
        elif line.startswith('!['):
            flush(); match = re.fullmatch(r'!\[([^]]*)\]\(([^)]+)\)', line)
            p = document.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(); run.add_picture(str((source.parent / match[2]).resolve()), width=Cm(13))
            props = run._r.xpath('.//wp:docPr')
            if props: props[0].set('descr', match[1])
            p.paragraph_format.keep_with_next = True
        elif line.startswith('- '):
            flush(); inline(document.add_paragraph(style='List Bullet'), line[2:])
        elif not line.strip():
            flush()
        else:
            pending.append(line.strip())
        i += 1
    flush()
    document.core_properties.title = 'Báo cáo thực hành TurtleBot 4: SLAM Toolbox, AMCL và Nav2'
    document.core_properties.subject = 'Mô phỏng, mapping, localization và autonomous navigation'
    document.core_properties.author = '[Điền họ và tên sinh viên]'
    document.core_properties.keywords = 'TurtleBot 4, ROS 2 Humble, SLAM Toolbox, AMCL, Nav2'
    document.core_properties.created = datetime.now(timezone.utc)
    document.core_properties.modified = datetime.now(timezone.utc)
    output = ROOT / 'docs/BAO_CAO_TURTLEBOT4.docx'
    document.save(output)
    print(output)


if __name__ == '__main__':
    main()
